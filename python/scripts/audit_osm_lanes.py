"""Audit corridor data quality and write ``docs/osm-lane-audit.md``.

Answers the questions the network file format design (ADR-0007) and the
validation plan depend on:

- How complete is OSM lane/turn/speed tagging on the v0.1 corridor?
- Which signals on the corridor exist, on what control system and mode?
- Which corridor intersections have turning movement counts, and how recent?
- Did the 2014-2017 Bluetooth travel-time program cover the corridor?

Inputs are produced by ``fetch_corridor_osm.py`` and
``fetch_toronto_open_data.py``; rerun those first (both cache into ``/data``).

Usage (venv with the ``data`` extra installed):
    python python/scripts/audit_osm_lanes.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import geopandas as gpd
import osmnx as ox
from openpyxl import load_workbook
from shapely.geometry import box

REPO_ROOT = Path(__file__).resolve().parents[2]
OSM_DIR = REPO_ROOT / "data" / "osm"
OPEN_DIR = REPO_ROOT / "data" / "open"
REPORT_PATH = REPO_ROOT / "docs" / "osm-lane-audit.md"

# Keep in sync with fetch_corridor_osm.py: (west, south, east, north).
BBOX = (-79.475, 43.650, -79.395, 43.745)

CORRIDOR_PATTERN = re.compile(
    r"eglinton|allen|st\.? clair|bathurst|dufferin|bloor", re.IGNORECASE
)
MAJOR_HIGHWAY_TYPES = {
    "motorway",
    "motorway_link",
    "trunk",
    "trunk_link",
    "primary",
    "primary_link",
    "secondary",
    "secondary_link",
    "tertiary",
    "tertiary_link",
}
TURN_LANE_COLUMNS = ["turn:lanes", "turn:lanes:forward", "turn:lanes:backward"]
CONTROL_MODE_HINTS = ("SA1", "SA2", "SAV", "PA-", "SCOOT", "CONTROL", "MODE")


def pct(numerator: int, denominator: int) -> str:
    return f"{100 * numerator / denominator:.0f}%" if denominator else "n/a"


def as_text(value: object) -> str:
    """GraphML round-trips merged edges as lists; normalize for matching."""
    if isinstance(value, list):
        return ";".join(str(v) for v in value)
    return "" if value is None else str(value)


def present(edges: gpd.GeoDataFrame, column: str) -> int:
    if column not in edges.columns:
        return 0
    return int(edges[column].notna().sum())


def any_turn_lanes(edges: gpd.GeoDataFrame) -> int:
    have = None
    for column in TURN_LANE_COLUMNS:
        if column not in edges.columns:
            continue
        mask = edges[column].notna()
        have = mask if have is None else (have | mask)
    return int(have.sum()) if have is not None else 0


def tag_coverage_rows(edges: gpd.GeoDataFrame, label: str) -> str:
    total = len(edges)
    return (
        f"| {label} | {total} | {pct(present(edges, 'lanes'), total)} "
        f"| {pct(any_turn_lanes(edges), total)} "
        f"| {pct(present(edges, 'maxspeed'), total)} |"
    )


def osm_sections(lines: list[str]) -> None:
    graph = ox.load_graphml(OSM_DIR / "corridor.graphml")
    nodes, edges = ox.convert.graph_to_gdfs(graph)

    highway_text = edges["highway"].map(as_text)
    is_major = highway_text.map(
        lambda h: any(part in MAJOR_HIGHWAY_TYPES for part in h.split(";"))
    )
    is_corridor = edges["name"].map(as_text).str.contains(CORRIDOR_PATTERN)

    lines += [
        "## OSM extract overview",
        "",
        f"- Drive network, bbox `{BBOX}` (west, south, east, north)",
        f"- {len(nodes):,} nodes, {len(edges):,} directed edges after simplification",
        f"- Signalized nodes (`highway=traffic_signals`): "
        f"{int((nodes.get('highway') == 'traffic_signals').sum()):,}",
        "",
        "### Way tag coverage (the lane-inference workload)",
        "",
        "| Edge set | Edges | `lanes` | `turn:lanes*` | `maxspeed` |",
        "|---|---|---|---|---|",
        tag_coverage_rows(edges, "All edges"),
        tag_coverage_rows(edges[is_major], "Major roads (motorway..tertiary)"),
        tag_coverage_rows(edges[is_corridor], "Corridor-named streets"),
        "",
        "### Corridor streets, individually",
        "",
        "| Street match | Edges | `lanes` | `turn:lanes*` | `maxspeed` |",
        "|---|---|---|---|---|",
    ]
    for street in (
        "Eglinton",
        "Allen",
        r"St\.? Clair",
        "Bathurst",
        "Dufferin",
        "Bloor",
    ):
        subset = edges[
            edges["name"].map(as_text).str.contains(street, case=False, regex=True)
        ]
        lines.append(tag_coverage_rows(subset, street.replace(r"\.?", "")))

    restrictions = json.loads(
        (OSM_DIR / "turn_restrictions.json").read_text(encoding="utf-8")
    )
    kinds = Counter(
        element.get("tags", {}).get("restriction", "<untagged>")
        for element in restrictions.get("elements", [])
    )
    lines += [
        "",
        "### Turn-restriction relations (Overpass, not in the osmnx graph)",
        "",
        f"- {sum(kinds.values())} relations in bbox: "
        + ", ".join(f"`{k}` x{v}" for k, v in kinds.most_common()),
        "",
    ]


def signals_section(lines: list[str]) -> None:
    west, south, east, north = BBOX
    corridor_signals = []
    with (OPEN_DIR / "traffic_signals.csv").open(encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            try:
                lon, lat = json.loads(row["geometry"])["coordinates"]
            except (KeyError, ValueError, TypeError):
                continue
            if not (west <= lon <= east and south <= lat <= north):
                continue
            streets = " / ".join(
                row.get(k) or ""
                for k in ("MAIN_STREET", "SIDE1_STREET", "SIDE2_STREET")
            )
            if CORRIDOR_PATTERN.search(streets):
                corridor_signals.append(row)

    modes = Counter(
        (row.get("SIGNALSYSTEM") or "?", row.get("CONTROL_MODE") or "?")
        for row in corridor_signals
    )
    lines += [
        "## Corridor traffic signals (Toronto Open Data: Traffic Signals Tabular)",
        "",
        f"- {len(corridor_signals)} signals in-bbox on corridor-named streets",
        "- System / control-mode distribution: "
        + ", ".join(f"`{sys_}/{mode}` x{n}" for (sys_, mode), n in modes.most_common()),
        "",
        "### Key corridor junctions",
        "",
        "| PX | Junction | System | Control mode | Transit preempt |",
        "|---|---|---|---|---|",
    ]
    for row in corridor_signals:
        main = row.get("MAIN_STREET") or ""
        side = row.get("SIDE1_STREET") or ""
        if CORRIDOR_PATTERN.search(main) and CORRIDOR_PATTERN.search(side):
            lines.append(
                f"| {row.get('PX')} | {main} / {side} "
                f"| {row.get('SIGNALSYSTEM')} | {row.get('CONTROL_MODE')} "
                f"| {row.get('TRANSIT_PREEMPT') or '-'} |"
            )
    lines.append("")


def control_mode_legend_section(lines: list[str]) -> None:
    workbook = load_workbook(
        OPEN_DIR / "traffic_signal_readme.xlsx", read_only=True, data_only=True
    )
    found: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None]
            text = " — ".join(cells)
            if cells and any(hint in text.upper() for hint in CONTROL_MODE_HINTS):
                found.append(f"- {text}")
    lines += [
        "## Control-mode legend (from the dataset's readme)",
        "",
        *(found[:30] or ["- (no control-mode definitions found in readme)"]),
        "",
    ]


def tmc_section(lines: list[str]) -> None:
    west, south, east, north = BBOX
    corridor_counts = []
    with (OPEN_DIR / "tmc_most_recent_summary.csv").open(
        encoding="utf-8-sig"
    ) as handle:
        for row in csv.DictReader(handle):
            try:
                lon, lat = float(row["longitude"]), float(row["latitude"])
            except (KeyError, ValueError):
                continue
            if not (west <= lon <= east and south <= lat <= north):
                continue
            if CORRIDOR_PATTERN.search(row.get("location_name") or ""):
                corridor_counts.append(row)

    years = Counter(
        (row.get("latest_count_date") or "????")[:4] for row in corridor_counts
    )
    lines += [
        "## Corridor turning movement counts (most recent per intersection)",
        "",
        f"- {len(corridor_counts)} counted corridor locations in-bbox",
        "- Latest-count year distribution: "
        + ", ".join(f"{year}: {n}" for year, n in sorted(years.items(), reverse=True)),
        "",
        "### Key corridor junctions",
        "",
        "| Location | Latest count | Duration (h) | Total vehicles |",
        "|---|---|---|---|",
    ]
    for row in sorted(corridor_counts, key=lambda r: r.get("latest_count_date") or ""):
        name = row.get("location_name") or ""
        if len(CORRIDOR_PATTERN.findall(name)) >= 2:
            lines.append(
                f"| {name} | {row.get('latest_count_date')} "
                f"| {row.get('count_duration')} | {row.get('total_vehicle')} |"
            )
    lines.append("")


def bluetooth_section(lines: list[str]) -> None:
    routes = gpd.read_file(f"zip://{OPEN_DIR / 'bluetooth_routes_wgs84.zip'}")
    corridor_poly = box(*BBOX)
    hits = routes[routes.geometry.intersects(corridor_poly)]
    name_column = next(
        (c for c in routes.columns if re.search(r"name|route|id", c, re.IGNORECASE)),
        routes.columns[0],
    )
    lines += [
        "## Bluetooth travel-time route coverage (2014-2017 program)",
        "",
        f"- {len(routes)} routes citywide; {len(hits)} intersect the corridor bbox",
        f"- Route attributes available: {', '.join(routes.columns)}",
        "- Corridor-intersecting routes: "
        + (", ".join(sorted(str(v) for v in hits[name_column])) or "(none)"),
        "",
    ]


def main() -> int:
    lines = [
        "# Corridor data audit — OSM tagging, signals, counts, travel times",
        "",
        f"Generated {date.today().isoformat()} by `python/scripts/audit_osm_lanes.py`",
        "from data fetched by `fetch_corridor_osm.py` and",
        "`fetch_toronto_open_data.py`. Regenerate by rerunning those scripts.",
        "Evidence base for ADR-0007 (network format) and `docs/validation.md`.",
        "",
    ]
    osm_sections(lines)
    signals_section(lines)
    control_mode_legend_section(lines)
    tmc_section(lines)
    bluetooth_section(lines)

    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {REPORT_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
